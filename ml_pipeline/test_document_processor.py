"""
Test for ML-002 (document_processor), enhanced with table extraction
and OCR fallback coverage.

Generates its own PDF/PPTX fixtures at runtime — no committed binaries.

Run:
    python -m ml_pipeline.test_document_processor
"""
import os
import tempfile

import pymupdf as fitz  # PyMuPDF's current import name
from pptx import Presentation

from ml_pipeline.document_processor import (
    process_document,
    process_document_structured,
    _PDFPLUMBER_AVAILABLE,
    _OCR_AVAILABLE,
)

SAMPLE_TEXT = "Stratified sampling divides the population into homogeneous subgroups."


def _make_fixture_pdf(dir_path: str) -> str:
    path = os.path.join(dir_path, "fixture.pdf")
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), SAMPLE_TEXT)
    doc.save(path)
    doc.close()
    return path


def _make_fixture_pptx(dir_path: str) -> str:
    path = os.path.join(dir_path, "fixture.pptx")
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Sampling Design"
    slide.placeholders[1].text = SAMPLE_TEXT
    prs.save(path)
    return path


def _make_empty_pptx(dir_path: str) -> str:
    path = os.path.join(dir_path, "empty.pptx")
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])  # blank layout, no text
    prs.save(path)
    return path


def _make_fixture_pdf_with_table(dir_path: str) -> str:
    """A PDF containing an actual drawn table (grid lines + cell text),
    since pdfplumber's table detection relies on visible structure, not
    just aligned whitespace."""
    path = os.path.join(dir_path, "fixture_table.pdf")
    doc = fitz.open()
    page = doc.new_page()

    rows, cols = 3, 2
    x0, y0, cell_w, cell_h = 72, 72, 100, 24
    data = [["Region", "Population"], ["North", "1200"], ["South", "980"]]

    for r in range(rows):
        for c in range(cols):
            x = x0 + c * cell_w
            y = y0 + r * cell_h
            rect = fitz.Rect(x, y, x + cell_w, y + cell_h)
            page.draw_rect(rect, color=(0, 0, 0), width=1)
            page.insert_text((x + 4, y + 16), data[r][c], fontsize=10)

    doc.save(path)
    doc.close()
    return path


def _make_fixture_scanned_pdf(dir_path: str) -> str:
    """A PDF with NO extractable text layer, only a rasterized image
    containing text — simulates a scanned document for the OCR path."""
    from PIL import Image, ImageDraw

    img_path = os.path.join(dir_path, "scanned_source.png")
    img = Image.new("RGB", (600, 200), color="white")
    draw = ImageDraw.Draw(img)
    draw.text((20, 80), "Scanned competency evidence", fill="black")
    img.save(img_path)

    path = os.path.join(dir_path, "fixture_scanned.pdf")
    doc = fitz.open()
    page = doc.new_page()
    rect = fitz.Rect(0, 0, page.rect.width, page.rect.height)
    page.insert_image(rect, filename=img_path)
    doc.save(path)
    doc.close()
    return path


def run_tests() -> bool:
    all_passed = True
    with tempfile.TemporaryDirectory() as tmp:
        # --- Regression: existing behavior must be unchanged ---
        pdf_path = _make_fixture_pdf(tmp)
        pdf_text = process_document(pdf_path)
        ok = "Stratified sampling" in pdf_text
        print(f"[{'PASS' if ok else 'FAIL'}] (regression) PDF text extraction: {pdf_text[:60]!r}")
        all_passed &= ok

        pptx_path = _make_fixture_pptx(tmp)
        pptx_text = process_document(pptx_path)
        ok = "Stratified sampling" in pptx_text and "Sampling Design" in pptx_text
        print(f"[{'PASS' if ok else 'FAIL'}] (regression) PPTX text extraction: {pptx_text[:60]!r}")
        all_passed &= ok

        empty_path = _make_empty_pptx(tmp)
        empty_text = process_document(empty_path)
        ok = empty_text == ""
        print(f"[{'PASS' if ok else 'FAIL'}] (regression) Empty PPTX returns '' : {empty_text!r}")
        all_passed &= ok

        try:
            process_document(os.path.join(tmp, "does_not_exist.pdf"))
            ok = False
        except FileNotFoundError:
            ok = True
        print(f"[{'PASS' if ok else 'FAIL'}] (regression) Missing file raises FileNotFoundError")
        all_passed &= ok

        bad_path = os.path.join(tmp, "notes.txt")
        with open(bad_path, "w") as f:
            f.write("plain text file")
        try:
            process_document(bad_path)
            ok = False
        except ValueError:
            ok = True
        print(f"[{'PASS' if ok else 'FAIL'}] (regression) Unsupported extension raises ValueError")
        all_passed &= ok

        # --- New: table extraction ---
        if _PDFPLUMBER_AVAILABLE:
            table_pdf_path = _make_fixture_pdf_with_table(tmp)
            structured = process_document_structured(table_pdf_path)
            tables_found = structured[0]["tables"]
            ok = len(tables_found) >= 1 and any(
                "Region" in cell for row in tables_found[0] for cell in row
            )
            print(f"[{'PASS' if ok else 'FAIL'}] table extraction finds the drawn table: {tables_found}")
            all_passed &= ok

            # Table content should also surface in the plain-text contract
            plain_text = process_document(table_pdf_path)
            ok = "[TABLE]" in plain_text and "Region" in plain_text
            print(f"[{'PASS' if ok else 'FAIL'}] table content appears in plain-text output")
            all_passed &= ok
        else:
            print("[SKIP] table extraction tests — pdfplumber not installed")

        # --- New: OCR fallback on a scanned (no text layer) PDF ---
        if _OCR_AVAILABLE:
            scanned_path = _make_fixture_scanned_pdf(tmp)
            structured = process_document_structured(scanned_path)
            page0 = structured[0]
            ok = page0["ocr_used"] is True and page0["content_type"] == "ocr"
            print(f"[{'PASS' if ok else 'FAIL'}] OCR fallback triggers on scanned page: content_type={page0['content_type']!r}")
            all_passed &= ok
            # Don't assert exact OCR text match — OCR accuracy varies by
            # environment/font rendering. Just confirm *something* came back.
            ok = len(page0["text"]) > 0
            print(f"[{'PASS' if ok else 'FAIL'}] OCR fallback extracted non-empty text: {page0['text'][:60]!r}")
            all_passed &= ok
        else:
            print("[SKIP] OCR fallback tests — pytesseract/Tesseract not available")

        # --- New: normal text page does NOT trigger OCR (cost control) ---
        structured = process_document_structured(pdf_path)
        ok = structured[0]["ocr_used"] is False and structured[0]["content_type"] == "text"
        print(f"[{'PASS' if ok else 'FAIL'}] normal text page does not trigger OCR")
        all_passed &= ok

        # --- New: graceful degradation if a page can't be table-parsed ---
        # (Using the plain fixture PDF, which has no table structure —
        # extraction must return [] for tables, not raise.)
        structured = process_document_structured(pdf_path)
        ok = structured[0]["tables"] == []
        print(f"[{'PASS' if ok else 'FAIL'}] page with no table returns empty tables list, no crash")
        all_passed &= ok

    return all_passed


if __name__ == "__main__":
    import sys
    success = run_tests()
    print("\nML-002 TEST SUITE:", "PASS" if success else "FAIL")
    sys.exit(0 if success else 1)
