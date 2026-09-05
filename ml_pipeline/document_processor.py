"""
ML-002 — Document extraction: PDF and PPTX -> plain text.

Scope is deliberately narrow: extraction only. No MCQ generation,
validation, embeddings, RAG, concept extraction, or adaptive selection
belongs in this module.
"""
from __future__ import annotations

import os

import pymupdf as fitz  # PyMuPDF's current import name
from pptx import Presentation

SUPPORTED_EXTENSIONS = {".pdf", ".pptx"}


def _extract_pdf_text(file_path: str) -> str:
    """Extract and concatenate text from every page of a PDF."""
    doc = fitz.open(file_path)
    try:
        pages = [page.get_text() for page in doc]
    finally:
        doc.close()
    return "\n".join(p for p in pages if p.strip())


def _extract_pptx_text(file_path: str) -> str:
    """Extract and concatenate text from every text-bearing shape on every slide."""
    presentation = Presentation(file_path)
    slide_texts = []
    for slide in presentation.slides:
        shape_texts = [
            shape.text for shape in slide.shapes
            if hasattr(shape, "text") and shape.text.strip()
        ]
        if shape_texts:
            slide_texts.append("\n".join(shape_texts))
    return "\n".join(slide_texts)


def process_document(file_path: str) -> str:
    """
    Extract plain text from a PDF or PPTX file.

    Args:
        file_path: path to a .pdf or .pptx file.

    Returns:
        Extracted text as a single string. Returns "" (empty string,
        not an error) if the file is valid but contains no extractable
        text — an empty document is a legitimate, expected case, not a
        failure.

    Raises:
        FileNotFoundError: if file_path does not exist.
        ValueError: if the file extension isn't .pdf or .pptx.
    """
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"process_document: no such file: {file_path}")

    ext = os.path.splitext(file_path)[1].lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            f"process_document: unsupported file type {ext!r} "
            f"(supported: {sorted(SUPPORTED_EXTENSIONS)})"
        )

    if ext == ".pdf":
        return _extract_pdf_text(file_path)
    return _extract_pptx_text(file_path)


if __name__ == "__main__":
    import sys
    if len(sys.argv) != 2:
        print("Usage: python ml_pipeline/document_processor.py <path-to-pdf-or-pptx>")
        sys.exit(1)
    text = process_document(sys.argv[1])
    print(f"Extracted {len(text)} characters.")
    print(text[:300])
